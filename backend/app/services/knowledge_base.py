"""RAG Knowledge Base - document ingestion, persistent storage, and retrieval.

Supports: PDF, Markdown, TXT, Code files.
Flow: File -> Chunk -> Embedding -> VectorStore -> Retriever -> LLM.

Persistence: when a SQLAlchemy session is passed (db != None), documents and
chunks are stored in knowledge_documents/knowledge_chunks and are read back from
there for every query. Without db, the process-wide in-memory index is the only
store, so unit tests and CLI use keep working.

Memory boundary: persisted content is never copied into the process-wide index.
A DB-backed query re-reads and re-embeds its own rows, so retaining them here
would only make the singleton grow with every account that touches the
knowledge base. Only the shared term vocabulary (bounded by
``SimpleEmbedder.MAX_VOCAB``) is warmed from persisted rows.
"""
import asyncio
import hashlib
import json
import logging
from pathlib import Path

import numpy as np
from core.text_tokens import iter_terms
from models.records import (
    KnowledgeChunk,
    KnowledgeCollection,
    KnowledgeCollectionDocument,
    KnowledgeDocument,
)

logger = logging.getLogger(__name__)


class SimpleEmbedder:
    """Lightweight embedding using TF-IDF-like bag-of-words vectors."""

    # Bounds process memory for long-running servers that ingest many files.
    MAX_VOCAB = 50_000

    def __init__(self):
        self.vocab: dict[str, int] = {}
        # Diagnostics: terms that could not be indexed because the vocabulary
        # hit MAX_VOCAB. They stay unretrievable, so this must not be silent.
        self.dropped_terms = 0

    def fit(self, texts: list[str]):
        capped = False
        for text in texts:
            for token in self._tokenize(text):
                if token in self.vocab:
                    continue
                if len(self.vocab) >= self.MAX_VOCAB:
                    self.dropped_terms += 1
                    capped = True
                    continue
                self.vocab[token] = len(self.vocab)
        if capped:
            logger.warning(
                "knowledge vocabulary reached the %d-term cap; %d term(s) were not "
                "indexed and documents made only of them cannot be retrieved",
                self.MAX_VOCAB,
                self.dropped_terms,
            )

    @property
    def capped(self) -> bool:
        """True once the vocabulary reached MAX_VOCAB and started dropping terms."""
        return self.dropped_terms > 0

    def embed(self, text: str) -> np.ndarray:
        vec = np.zeros(len(self.vocab) or 1, dtype=np.float32)
        tokens = self._tokenize(text)
        for token in tokens:
            idx = self.vocab.get(token)
            if idx is not None:
                vec[idx] += 1.0
        norm = np.linalg.norm(vec)
        if norm > 0:
            vec /= norm
        return vec

    def embed_batch(self, texts: list[str]) -> list[np.ndarray]:
        return [self.embed(t) for t in texts]

    def _tokenize(self, text: str) -> list[str]:
        # CJK runs are not word-segmented by \w; bigrams keep Chinese queries
        # searchable (see core.text_tokens).
        return list(iter_terms(text))


class InMemoryVectorStore:
    """In-memory vector store with cosine similarity search."""

    def __init__(self):
        self.documents: list[dict] = []
        self.vectors: list[np.ndarray] = []

    def add(self, doc_id: str, text: str, metadata: dict, vector: np.ndarray):
        self.documents.append({"id": doc_id, "text": text, "metadata": metadata})
        self.vectors.append(vector)

    def search(self, query_vector: np.ndarray, top_k: int = 5) -> list[dict]:
        if not self.vectors:
            return []
        scores = np.array([np.dot(query_vector, v) for v in self.vectors])
        top_indices = np.argsort(scores)[::-1][:top_k]
        results = []
        for idx in top_indices:
            if scores[idx] > 0:
                doc = self.documents[idx].copy()
                doc["score"] = float(scores[idx])
                results.append(doc)
        return results

    def clear(self):
        self.documents = []
        self.vectors = []

    def remove_by_metadata(self, key: str, value):
        keep = [
            (d, v) for d, v in zip(self.documents, self.vectors, strict=False)
            if d["metadata"].get(key) != value
        ]
        self.documents = [d for d, _ in keep]
        self.vectors = [v for _, v in keep]


class TextChunker:
    """Splits text into overlapping chunks."""

    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.chunk_size = max(1, int(chunk_size))
        # An overlap of chunk_size or more leaves no room for the next pass to
        # advance, so it is clamped below the chunk size.
        self.chunk_overlap = max(0, min(int(chunk_overlap), self.chunk_size - 1))

    def split(self, text: str) -> list[str]:
        paragraphs = text.split("\n\n")
        chunks = []
        current = ""
        for para in paragraphs:
            para = para.strip()
            if not para:
                continue
            separator = "\n\n" if current else ""
            # The separator counts towards the ceiling: merging two paragraphs
            # used to produce chunks two characters above the configured size.
            if len(current) + len(separator) + len(para) <= self.chunk_size:
                current += separator + para
            else:
                if current:
                    chunks.append(current)
                current = para
                while len(current) > self.chunk_size:
                    split_point = current[:self.chunk_size].rfind(" ")
                    # Every pass must leave at least one new character for the
                    # next one. A break inside the overlap window (or no space
                    # at all, as in long CJK runs) used to slice ``current``
                    # with the same offset forever, hanging the whole process.
                    if split_point < self.chunk_overlap + 1:
                        split_point = self.chunk_size
                    if current[:split_point].strip():
                        chunks.append(current[:split_point])
                    current = current[split_point - self.chunk_overlap:]
        if current.strip():
            chunks.append(current)
        return chunks


class FileParser:
    """Parse various file types into plain text."""

    SUPPORTED_EXTENSIONS = {
        ".txt", ".md", ".markdown", ".py", ".js", ".ts", ".java",
        ".go", ".rs", ".cpp", ".c", ".h", ".html", ".css", ".json",
        ".yaml", ".yml", ".xml", ".toml", ".cfg", ".ini",
        # V1.4: spreadsheets and Word documents are first-class sources.
        ".csv", ".tsv", ".docx",
    }

    def parse(self, filepath: str) -> tuple[str, dict]:
        path = Path(filepath)
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf(path)
        if ext == ".docx":
            return self._parse_docx(path)
        if ext in {".csv", ".tsv"}:
            return self._parse_table(path, delimiter="," if ext == ".csv" else "\t")
        elif ext in self.SUPPORTED_EXTENSIONS:
            return self._parse_text(path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _parse_text(self, path: Path) -> tuple[str, dict]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return content, {"filename": path.name, "type": "text", "size": path.stat().st_size}

    def _parse_pdf(self, path: Path) -> tuple[str, dict]:
        try:
            import PyPDF2
            text = ""
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text() or ""
            return text.strip(), {"filename": path.name, "type": "pdf", "pages": len(reader.pages)}
        except ImportError:
            pass
        try:
            import pdfplumber
            text = ""
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    text += page.extract_text() or ""
            return text.strip(), {"filename": path.name, "type": "pdf"}
        except ImportError:
            raise RuntimeError(
                "PDF parsing requires PyPDF2 or pdfplumber. Install with: pip install PyPDF2"
            )

    def _parse_table(self, path: Path, *, delimiter: str) -> tuple[str, dict]:
        """Flatten a CSV/TSV into ``header: value`` lines the chunker can split."""
        import csv

        rows: list[str] = []
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            header: list[str] = []
            for index, row in enumerate(reader):
                if index == 0:
                    header = [cell.strip() for cell in row]
                    continue
                cells = [cell.strip() for cell in row]
                if not any(cells):
                    continue
                if header and len(header) == len(cells):
                    rows.append("\n".join(f"{name}: {value}" for name, value in zip(header, cells, strict=False)))
                else:
                    rows.append(" | ".join(cells))
        text = "\n\n".join(rows)
        return text, {"filename": path.name, "type": "table", "rows": len(rows), "columns": header}

    def _parse_docx(self, path: Path) -> tuple[str, dict]:
        """Extract paragraph + table text from a .docx package.

        ``python-docx`` is used when installed; otherwise the file is read as an
        OOXML zip, which needs no extra dependency and still yields the text.
        """
        try:
            import docx  # type: ignore

            document = docx.Document(str(path))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            return "\n\n".join(parts), {"filename": path.name, "type": "docx", "paragraphs": len(parts)}
        except ImportError:
            pass
        import re
        import zipfile

        try:
            with zipfile.ZipFile(path) as archive:
                xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
        except (KeyError, OSError, zipfile.BadZipFile) as exc:
            raise RuntimeError("无法读取 DOCX 文件（不是有效的 Word 文档）") from exc
        # Paragraph boundaries first, then strip tags so text does not run together.
        xml = re.sub(r"</w:p>", "\n", xml)
        text = re.sub(r"<[^>]+>", "", xml)
        text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
        return text, {"filename": path.name, "type": "docx", "paragraphs": len(text.splitlines())}


#: Retrieval strategies supported by the knowledge base (V1.4).
RETRIEVAL_MODES = ("semantic", "keyword", "hybrid", "rerank")


def normalize_retrieval_mode(value: str | None) -> str:
    mode = str(value or "semantic").strip().lower()
    if mode not in RETRIEVAL_MODES:
        raise ValueError("retrieval mode must be one of " + ", ".join(RETRIEVAL_MODES))
    return mode


def lexical_score(question: str, text: str) -> float:
    """Term-overlap score in ``[0, 1]`` used by keyword/hybrid/rerank modes."""
    query_terms = list(dict.fromkeys(iter_terms(question)))
    if not query_terms:
        return 0.0
    haystack = set(iter_terms(text))
    if not haystack:
        return 0.0
    hits = sum(1 for term in query_terms if term in haystack)
    return hits / len(query_terms)


def _combine_scores(mode: str, semantic: float, lexical: float) -> float:
    if mode == "keyword":
        return lexical
    if mode == "hybrid" or mode == "rerank":
        # Cosine similarity and term overlap are both in [0, 1]; the weights
        # keep semantics dominant while still rewarding exact term hits.
        return 0.6 * semantic + 0.4 * lexical
    return semantic


def _rank_scored(scored: list[tuple], mode: str, top_k: int) -> list[tuple]:
    """Rank ``(payload, semantic, lexical)`` rows for one retrieval mode."""
    if mode == "rerank":
        # Recall with the vector score, then re-order the shortlist lexically.
        shortlist = sorted(scored, key=lambda item: item[1], reverse=True)[: max(top_k * 3, top_k)]
        ordered = sorted(
            shortlist, key=lambda item: _combine_scores(mode, item[1], item[2]), reverse=True
        )[:top_k]
    else:
        ordered = sorted(
            scored, key=lambda item: _combine_scores(mode, item[1], item[2]), reverse=True
        )[:top_k]
    return [
        (item[0], _combine_scores(mode, item[1], item[2]), item[1], item[2]) for item in ordered
    ]


class KnowledgeBase:
    """RAG Knowledge Base: ingest documents and query them (optionally persistent)."""

    def __init__(self):
        self.vector_store = InMemoryVectorStore()
        self.embedder = SimpleEmbedder()
        self.chunker = TextChunker()
        self.parser = FileParser()
        # One entry per owner whose vocabulary was warmed in this process.
        self._loaded_scopes: set[int | None] = set()

    def _vector_width(self) -> int:
        return len(self.embedder.vocab) or 1

    def _count_documents(self) -> int:
        """Distinct local documents currently held in the in-memory index."""
        return len(
            {doc["metadata"].get("filename") for doc in self.vector_store.documents}
        )

    @staticmethod
    def _pad_vector(vector, width: int) -> np.ndarray:
        """Zero-pad an older vector after the shared vocabulary grew."""
        if vector is None:
            return np.zeros(width, dtype=np.float32)
        if len(vector) == width:
            return vector
        padded = np.zeros(width, dtype=np.float32)
        padded[: len(vector)] = vector
        return padded

    def _ensure_loaded(self, db=None, user_id: int | None = None):
        """Warm the shared vocabulary from the caller's persisted rows (once per owner).

        Persisted content is deliberately *not* copied into the process-wide
        in-memory index. A DB-backed query re-reads and re-embeds its own rows,
        so keeping a second copy here only made the singleton grow with every
        account that touched the knowledge base. Warming the vocabulary is the
        part that cannot be re-derived per query: it is what makes a question
        comparable with the stored rows.
        """
        if db is None or user_id in self._loaded_scopes:
            return
        try:
            chunks_query = (
                db.query(KnowledgeChunk)
                .join(KnowledgeDocument, KnowledgeChunk.doc_id == KnowledgeDocument.id)
                .order_by(KnowledgeChunk.doc_id, KnowledgeChunk.chunk_index)
            )
            if user_id is not None:
                chunks_query = chunks_query.filter(KnowledgeDocument.user_id == user_id)
            self.embedder.fit([ch.content for ch in chunks_query.all()])
        except Exception:
            # DB not ready (e.g. table missing) -> in-memory only. The request
            # still reports its own error; here we only lose retrieval quality.
            logger.warning(
                "knowledge vocabulary warm-up failed for user_id=%s", user_id,
                exc_info=True,
            )
        self._loaded_scopes.add(user_id)

    def upload(self, filepath: str, db=None, user_id: int | None = None, filename: str | None = None) -> dict:
        """Ingest a file: parse, chunk, embed, index (and persist when db given)."""
        text, metadata = self.parser.parse(filepath)
        if filename:
            metadata["filename"] = filename
        if not text.strip():
            return {"status": "empty", "file": filepath, "chunks": 0}

        chunks = self.chunker.split(text)
        document_name = metadata.get("filename", filename or Path(filepath).name)
        # One filename is one document: re-uploading replaces the previous copy.
        # Keeping both made chunks()/query() read the stale copy and made
        # delete_document() remove only one of them.
        if db is not None and user_id is not None:
            self._delete_documents_by_name(db, user_id, document_name)
        self._ensure_loaded(db, user_id)
        self.embedder.fit(chunks)
        chunk_meta_base = {"filename": document_name, "type": metadata.get("type", "text")}
        # The process-wide in-memory index only backs the local (session-less)
        # mode. With a session the database is the single source of truth, and a
        # second copy here would keep every uploader's text in memory forever.
        if db is None:
            self.vector_store.remove_by_metadata("filename", document_name)
            # New tokens only append to the vocabulary, so existing vectors keep
            # their dot products once padded to the new width. Re-embedding the
            # whole corpus on every upload made large libraries progressively
            # slower for no gain.
            width = self._vector_width()
            self.vector_store.vectors = [
                self._pad_vector(vector, width) for vector in self.vector_store.vectors
            ]
            new_vecs = self.embedder.embed_batch(chunks)
            file_id = hashlib.md5(filepath.encode()).hexdigest()[:12]
            for i, (chunk, vector) in enumerate(zip(chunks, new_vecs, strict=False)):
                doc_id = f"{file_id}_{i}"
                chunk_meta = {**chunk_meta_base, "chunk_index": i, "total_chunks": len(chunks)}
                self.vector_store.add(doc_id, chunk, chunk_meta, vector)

        if db is not None and user_id is not None:
            doc = KnowledgeDocument(
                user_id=user_id,
                filename=document_name,
                filetype=metadata.get("type", "text"),
                chunk_count=len(chunks),
                doc_meta=json.dumps(metadata, ensure_ascii=False),
            )
            db.add(doc)
            db.flush()
            for i, chunk in enumerate(chunks):
                db.add(
                    KnowledgeChunk(
                        doc_id=doc.id,
                        chunk_index=i,
                        content=chunk,
                        meta=json.dumps({**chunk_meta_base, "chunk_index": i, "total_chunks": len(chunks)}, ensure_ascii=False),
                    )
                )
            db.commit()

        return {
            "status": "ingested",
            "file": filepath,
            "chunks": len(chunks),
            "type": metadata.get("type", "unknown"),
        }

    @staticmethod
    def _delete_documents_by_name(db, user_id: int, filename: str) -> int:
        """Remove every stored copy of ``filename`` owned by ``user_id``."""
        docs = (
            db.query(KnowledgeDocument)
            .filter(KnowledgeDocument.user_id == user_id, KnowledgeDocument.filename == filename)
            .all()
        )
        for doc in docs:
            db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == doc.id).delete()
            db.delete(doc)
        if docs:
            db.flush()
        return len(docs)

    @staticmethod
    def normalize_binding(binding: dict | None = None) -> dict:
        """Normalize an explicit knowledge range without broadening access."""
        raw = dict(binding or {})
        collection_ids = list(dict.fromkeys(str(item) for item in raw.get("collection_ids") or raw.get("collections") or [] if item))
        mode = str(raw.get("mode") or ("collections" if collection_ids else "all"))
        if mode not in {"all", "collections", "disabled"}:
            raise ValueError("knowledge binding mode must be all, collections, or disabled")
        if mode == "collections" and not collection_ids:
            raise ValueError("collections knowledge binding requires collection_ids")
        return {"mode": mode, "collection_ids": collection_ids}

    def query(
        self,
        question: str,
        top_k: int = 5,
        db=None,
        user_id: int | None = None,
        knowledge_binding: dict | None = None,
        retrieval_mode: str | None = None,
    ) -> dict:
        self._ensure_loaded(db, user_id)
        binding = self.normalize_binding(knowledge_binding)
        mode = normalize_retrieval_mode(retrieval_mode)
        if binding["mode"] == "disabled":
            return {
                "question": question,
                "results": [],
                "total_results": 0,
                "knowledge_binding": binding,
                "retrieval_mode": mode,
            }
        query_vector = self.embedder.embed(question)
        if not query_vector.any() and self.embedder.vocab:
            # No indexed term at all: every score would be zero and the caller
            # would silently be told the knowledge base has nothing relevant.
            logger.warning(
                "knowledge question matched no indexed term (vocab_size=%d, dropped_terms=%d)",
                len(self.embedder.vocab),
                self.embedder.dropped_terms,
            )
        if db is not None and user_id is not None:
            query = (
                db.query(KnowledgeChunk)
                .join(KnowledgeDocument, KnowledgeChunk.doc_id == KnowledgeDocument.id)
                .filter(KnowledgeDocument.user_id == user_id)
            )
            if binding["mode"] == "collections":
                query = (
                    query.join(KnowledgeCollectionDocument, KnowledgeCollectionDocument.document_id == KnowledgeDocument.id)
                    .join(KnowledgeCollection, KnowledgeCollection.id == KnowledgeCollectionDocument.collection_id)
                    .filter(KnowledgeCollection.user_id == user_id, KnowledgeCollection.id.in_(binding["collection_ids"]))
                )
            rows = query.distinct().all()
            collection_names: dict[int, list[dict]] = {}
            if rows:
                doc_ids = {row.doc_id for row in rows}
                memberships = (
                    db.query(KnowledgeCollectionDocument.document_id, KnowledgeCollection.id, KnowledgeCollection.name)
                    .join(KnowledgeCollection, KnowledgeCollection.id == KnowledgeCollectionDocument.collection_id)
                    .filter(KnowledgeCollectionDocument.document_id.in_(doc_ids), KnowledgeCollection.user_id == user_id)
                    .all()
                )
                for document_id, collection_id, name in memberships:
                    collection_names.setdefault(document_id, []).append({"id": collection_id, "name": name})
            vectors = self.embedder.embed_batch([row.content for row in rows]) if rows else []
            scored = [
                (row, float(np.dot(query_vector, vector)), lexical_score(question, row.content))
                for row, vector in zip(rows, vectors, strict=False)
            ]
            ranked = _rank_scored(scored, mode, top_k)
            results = [
                {
                    "text": row.content,
                    "score": round(score, 6),
                    "semantic_score": round(semantic, 6),
                    "lexical_score": round(lexical, 6),
                    "metadata": json.loads(row.meta) if row.meta else {},
                    "document_id": row.doc_id,
                    "chunk_id": row.id,
                    "collections": collection_names.get(row.doc_id, []),
                }
                for row, score, semantic, lexical in ranked
                if score > 0
            ]
        else:
            memory_hits = self.vector_store.search(query_vector, top_k=max(top_k, top_k * 3))
            scored = [
                (item, float(item["score"]), lexical_score(question, item["text"]))
                for item in memory_hits
            ]
            ranked = _rank_scored(scored, mode, top_k)
            results = []
            for item, score, semantic, lexical in ranked:
                if score <= 0:
                    continue
                enriched = dict(item)
                enriched["score"] = round(score, 6)
                enriched["semantic_score"] = round(semantic, 6)
                enriched["lexical_score"] = round(lexical, 6)
                results.append(enriched)
        return {
            "question": question,
            "results": [
                {
                    "text": item["text"][:300],
                    "score": round(float(item["score"]), 4),
                    "source": item["metadata"].get("filename", ""),
                    "chunk_index": item["metadata"].get("chunk_index"),
                    "document_id": item.get("document_id"),
                    "chunk_id": item.get("chunk_id"),
                    "collections": item.get("collections", []),
                    "semantic_score": item.get("semantic_score"),
                    "lexical_score": item.get("lexical_score"),
                }
                for item in results
            ],
            "total_results": len(results),
            "knowledge_binding": binding,
            "retrieval_mode": mode,
        }

    async def answer(
        self, question: str, top_k: int = 5, db=None, user_id: int | None = None, runtime=None, model: str = "default-model", knowledge_binding: dict | None = None, retrieval_mode: str | None = None
    ) -> dict:
        """RAG answer: retrieve relevant chunks, then generate with the runtime."""
        # Retrieval is synchronous (DB read + embedding), so it must not block
        # the event loop: a large library used to freeze every other request in
        # the process for the whole duration of the query.
        query_result = await asyncio.to_thread(
            self.query,
            question,
            top_k=top_k,
            db=db,
            user_id=user_id,
            knowledge_binding=knowledge_binding,
            retrieval_mode=retrieval_mode,
        )
        sources = query_result["results"]
        if not sources:
            return {"answer": "知识库中没有找到相关内容。", "sources": []}
        context = "\n\n".join(f"[{s['source']}] {s['text']}" for s in sources)
        prompt = f"[知识库内容]\n{context}\n\n问题: {question}\n请基于以上知识库内容回答。"
        if runtime is None:
            return {
                "answer": "未配置运行时，无法生成回答。检索结果如下：\n" + context,
                "sources": sources,
            }
        result = await runtime.chat(model, [{"role": "user", "content": prompt}])
        return {"answer": result.get("content", ""), "sources": sources}

    def documents(self, db=None, user_id: int | None = None) -> list[dict]:
        if db is not None:
            query = db.query(KnowledgeDocument)
            if user_id is not None:
                query = query.filter(KnowledgeDocument.user_id == user_id)
            docs = query.order_by(KnowledgeDocument.created_at.desc()).all()
            return [d.to_dict() for d in docs]
        seen = {}
        for doc in self.vector_store.documents:
            name = doc["metadata"].get("filename", "?")
            seen.setdefault(name, {"filename": name, "chunks": 0})
            seen[name]["chunks"] += 1
        return list(seen.values())

    def delete_document(self, filename: str, db=None, user_id: int | None = None) -> bool:
        removed = 0
        if db is not None:
            if user_id is not None:
                removed = self._delete_documents_by_name(db, user_id, filename)
            else:
                doc = db.query(KnowledgeDocument).filter(KnowledgeDocument.filename == filename).first()
                if doc is not None:
                    db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == doc.id).delete()
                    db.delete(doc)
                    removed = 1
            db.commit()
        before = len(self.vector_store.documents)
        self.vector_store.remove_by_metadata("filename", filename)
        return removed > 0 or len(self.vector_store.documents) < before

    def chunks(self, filename: str, db=None, user_id: int | None = None) -> list[dict]:
        if db is not None:
            query = db.query(KnowledgeDocument).filter(KnowledgeDocument.filename == filename)
            if user_id is not None:
                query = query.filter(KnowledgeDocument.user_id == user_id)
            # Legacy rows could hold several copies of one name; the newest one wins.
            doc = query.order_by(KnowledgeDocument.created_at.desc(), KnowledgeDocument.id.desc()).first()
            if doc is None:
                return []
            rows = (
                db.query(KnowledgeChunk)
                .filter(KnowledgeChunk.doc_id == doc.id)
                .order_by(KnowledgeChunk.chunk_index)
                .all()
            )
            return [r.to_dict() for r in rows]
        return [
            {"chunk_index": d["metadata"].get("chunk_index"), "content": d["text"][:300]}
            for d in self.vector_store.documents
            if d["metadata"].get("filename") == filename
        ]

    def stats(self, db=None, user_id: int | None = None) -> dict:
        self._ensure_loaded(db, user_id)
        diagnostics = {
            "vocab_size": len(self.embedder.vocab),
            "vocab_capped": self.embedder.capped,
            "vocab_dropped_terms": self.embedder.dropped_terms,
        }
        if db is not None and user_id is not None:
            docs = db.query(KnowledgeDocument).filter(KnowledgeDocument.user_id == user_id).all()
            return {
                "documents": len(docs),
                "chunks": sum(doc.chunk_count or 0 for doc in docs),
                **diagnostics,
            }
        return {
            "documents": self._count_documents(),
            "chunks": len(self.vector_store.documents),
            **diagnostics,
        }


_kb = None

def get_global_kb() -> KnowledgeBase:
    """Process-wide singleton knowledge base."""
    global _kb
    if _kb is None:
        _kb = KnowledgeBase()
    return _kb
